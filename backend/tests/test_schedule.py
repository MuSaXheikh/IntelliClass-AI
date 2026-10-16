"""Schedule: sheet upload, visibility, and the cancel-before-the-day rule."""

import io
from datetime import date, datetime, time, timedelta

import openpyxl
from fastapi.testclient import TestClient

from app.core.config import get_settings
from app.services.schedule.service import today_in
from tests.conftest import auth, register_login


def _xlsx(rows: list[list[object]], headers: list[str] | None = None) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(headers or ["course_code", "date", "start", "end", "topic", "location"])
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _upload(
    client: TestClient, token: str, data: bytes, name: str = "plan.xlsx", mode: str = "merge"
) -> dict:
    res = client.post(
        "/api/v1/schedule/upload",
        files={"file": (name, data)},
        data={"mode": mode},
        headers=auth(token),
    )
    assert res.status_code == 200, res.text
    return res.json()


def test_upload_is_all_or_nothing_and_visible_to_enrolled_students(
    client: TestClient, live_session: dict, instructor: tuple[str, dict], student: tuple[str, dict]
) -> None:
    itoken, _ = instructor
    today = today_in(get_settings())
    d1, d2 = today + timedelta(days=20), today + timedelta(days=22)
    bad = _xlsx(
        [
            [
                "CS-401",
                datetime(d1.year, d1.month, d1.day),
                time(9, 0),
                time(10, 30),
                "Intro",
                "Room 1",
            ],
            ["XX-999", d2.strftime("%d/%m/%Y"), "9:00 AM", "", "Ghost", ""],
        ]
    )
    result = _upload(client, itoken, bad)
    assert result["applied"] is False and "XX-999" in result["errors"][0]
    assert client.get("/api/v1/schedule", headers=auth(itoken)).json() == []

    good = _xlsx(
        [
            [
                "CS-401",
                datetime(d1.year, d1.month, d1.day),
                time(9, 0),
                time(10, 30),
                "Intro",
                "Room 1",
            ],
            ["cs-401", d2.strftime("%d/%m/%Y"), "9:00 AM", "", "Landmarks", ""],
        ]
    )
    result = _upload(client, itoken, good)
    assert result == {"applied": True, "added": 2, "updated": 0, "removed": 0, "errors": []}

    mine = client.get("/api/v1/schedule", headers=auth(itoken)).json()
    assert [e["scheduled_date"] for e in mine] == [d1.isoformat(), d2.isoformat()]
    assert mine[0]["start_time"] == "09:00:00" and mine[0]["class_name"] == "AI 101"
    assert mine[0]["can_cancel"] is True and mine[1]["end_time"] is None

    seen_by_student = client.get("/api/v1/schedule", headers=auth(student[0])).json()
    assert [e["id"] for e in seen_by_student] == [e["id"] for e in mine]
    outsider, _ = register_login(client, "student", "nobody@uni.edu", "mtf00000001")
    assert client.get("/api/v1/schedule", headers=auth(outsider)).json() == []
    other, _ = register_login(client, "instructor", "other@uni.edu")
    assert client.get("/api/v1/schedule", headers=auth(other)).json() == []

    # merge updates a row in place; replace drops upcoming rows missing from the sheet
    again = _xlsx([["CS-401", d1.isoformat(), "09:00", "11:00", "Intro (longer)", "Room 1"]])
    assert _upload(client, itoken, again)["updated"] == 1
    assert _upload(client, itoken, again, mode="replace")["removed"] == 1
    assert len(client.get("/api/v1/schedule", headers=auth(itoken)).json()) == 1


def test_csv_without_course_code_works_for_a_single_class(
    client: TestClient, live_session: dict, instructor: tuple[str, dict]
) -> None:
    itoken, _ = instructor
    day = (today_in(get_settings()) + timedelta(days=3)).isoformat()
    csv = f"date,start,topic\n{day},14:00,Revision\n".encode()
    assert _upload(client, itoken, csv, name="plan.csv")["added"] == 1
    past = b"date,start\n2020-01-01,09:00\n"
    result = _upload(client, itoken, past, name="plan.csv")
    assert result["applied"] is False and "in the past" in result["errors"][0]


def test_cancel_only_before_the_day(
    client: TestClient, live_session: dict, instructor: tuple[str, dict], student: tuple[str, dict]
) -> None:
    itoken, _ = instructor
    today = today_in(get_settings())
    later, same_day = today + timedelta(days=5), today
    _upload(
        client,
        itoken,
        _xlsx(
            [["CS-401", later.isoformat(), "09:00"], ["CS-401", same_day.isoformat(), "16:00"]],
            ["course_code", "date", "start"],
        ),
    )
    entries = {
        e["scheduled_date"]: e for e in client.get("/api/v1/schedule", headers=auth(itoken)).json()
    }
    assert entries[same_day.isoformat()]["can_cancel"] is False

    too_late = client.post(
        f"/api/v1/schedule/{entries[same_day.isoformat()]['id']}/cancel", headers=auth(itoken)
    )
    assert too_late.status_code == 409 and too_late.json()["error"]["code"] == "TOO_LATE_TO_CANCEL"

    future_id = entries[later.isoformat()]["id"]
    cancelled = client.post(
        f"/api/v1/schedule/{future_id}/cancel", json={"reason": "Conference"}, headers=auth(itoken)
    )
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "cancelled" and cancelled.json()["can_restore"] is True
    assert (
        client.post(f"/api/v1/schedule/{future_id}/cancel", headers=auth(itoken)).status_code == 409
    )
    for_student = client.get("/api/v1/schedule", headers=auth(student[0])).json()
    assert {e["status"] for e in for_student} == {"cancelled", "scheduled"}
    assert next(e for e in for_student if e["id"] == future_id)["cancel_reason"] == "Conference"

    assert (
        client.post(f"/api/v1/schedule/{future_id}/cancel", headers=auth(student[0])).status_code
        == 403
    )
    other, _ = register_login(client, "instructor", "other2@uni.edu")
    assert (
        client.post(f"/api/v1/schedule/{future_id}/restore", headers=auth(other)).status_code == 403
    )
    restored = client.post(f"/api/v1/schedule/{future_id}/restore", headers=auth(itoken))
    assert restored.status_code == 200 and restored.json()["status"] == "scheduled"


def test_template_download(
    client: TestClient, live_session: dict, instructor: tuple[str, dict]
) -> None:
    res = client.get("/api/v1/schedule/template.xlsx", headers=auth(instructor[0]))
    assert res.status_code == 200 and res.content[:2] == b"PK"
    sheet = openpyxl.load_workbook(io.BytesIO(res.content)).active
    assert [c.value for c in sheet[1]] == [
        "course_code",
        "date",
        "start",
        "end",
        "topic",
        "location",
    ]
    assert sheet["A2"].value == "CS-401" and date.fromisoformat(
        sheet["B2"].value
    ) > date.today() - timedelta(days=1)
    # the template uploads as it is, so a teacher can start from the examples
    assert _upload(client, instructor[0], res.content)["added"] == 2


def test_export_round_trips_through_upload(
    client: TestClient, live_session: dict, instructor: tuple[str, dict], student: tuple[str, dict]
) -> None:
    itoken, _ = instructor
    day = (today_in(get_settings()) + timedelta(days=9)).isoformat()
    _upload(client, itoken, _xlsx([["CS-401", day, "10:00", "11:00", "Exported", "Lab 2"]]))
    res = client.get("/api/v1/schedule/export.xlsx", headers=auth(itoken))
    assert res.status_code == 200 and res.content[:2] == b"PK"
    sheet = openpyxl.load_workbook(io.BytesIO(res.content)).active
    header = [c.value for c in sheet[1]]
    assert header[:6] == ["course_code", "date", "start", "end", "topic", "location"]
    assert [c.value for c in sheet[2]][:6] == ["CS-401", day, "10:00", "11:00", "Exported", "Lab 2"]
    # uploading the export again changes nothing
    assert _upload(client, itoken, res.content) == {
        "applied": True,
        "added": 0,
        "updated": 0,
        "removed": 0,
        "errors": [],
    }
    # students get their own view as a sheet too
    theirs = client.get("/api/v1/schedule/export.xlsx", headers=auth(student[0]))
    assert theirs.status_code == 200
    assert openpyxl.load_workbook(io.BytesIO(theirs.content)).active.max_row == 2


def test_spreadsheet_quirks_are_understood(
    client: TestClient, live_session: dict, instructor: tuple[str, dict]
) -> None:
    itoken, _ = instructor
    base = today_in(get_settings()) + timedelta(days=12)
    # a General-formatted time cell holds a fraction of a day; exported dates carry a time part
    sheet = _xlsx(
        [
            ["CS-401", f"{base.isoformat()} 00:00:00", 0.375, 0.5, "Serial time", ""],
            ["CS-401", (base + timedelta(days=1)).strftime("%d/%m/%Y 00:00"), "9:30", "", "", ""],
        ]
    )
    assert _upload(client, itoken, sheet)["added"] == 2
    times = {
        e["scheduled_date"]: e for e in client.get("/api/v1/schedule", headers=auth(itoken)).json()
    }
    assert times[base.isoformat()]["start_time"] == "09:00:00"
    assert times[base.isoformat()]["end_time"] == "12:00:00"
    assert times[(base + timedelta(days=1)).isoformat()]["start_time"] == "09:30:00"


def test_blank_and_explicit_course_code_for_the_same_slot_is_a_row_error(
    client: TestClient, live_session: dict, instructor: tuple[str, dict]
) -> None:
    itoken, _ = instructor
    day = (today_in(get_settings()) + timedelta(days=4)).isoformat()
    sheet = _xlsx([["", day, "09:00"], ["CS-401", day, "09:00"]], ["course_code", "date", "start"])
    result = _upload(client, itoken, sheet)
    assert result["applied"] is False and "Row 3: duplicate" in result["errors"][0]
    assert client.get("/api/v1/schedule", headers=auth(itoken)).json() == []
