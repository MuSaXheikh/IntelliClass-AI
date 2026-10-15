"""Deck conversion: type detection, PDF passthrough, and LibreOffice conversion when present."""

import io
import shutil
import zipfile
from pathlib import Path

import fitz
import pytest
from fastapi.testclient import TestClient

from app.core.errors import ApiError
from app.services.slides.convert import detect_kind, to_pdf
from tests.conftest import auth, register_login

HAS_SOFFICE = shutil.which("soffice") is not None


def _zip(entries: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    return buf.getvalue()


def _pptx(titles: list[str]) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for title in titles:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = f"Notes about {title.lower()}"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_detect_kind_from_bytes() -> None:
    assert detect_kind(b"%PDF-1.7 ...") == "pdf"
    assert detect_kind(_zip({"ppt/presentation.xml": b"<p/>"})) == "pptx"
    odp = _zip({"mimetype": b"application/vnd.oasis.opendocument.presentation"})
    assert detect_kind(odp) == "odp"
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 32
    assert detect_kind(ole, "lecture.PPT") == "ppt"
    assert detect_kind(ole, "lecture.doc") is None
    assert detect_kind(_zip({"word/document.xml": b"<w/>"})) is None
    assert detect_kind(b"PK\x03\x04 not really a zip") is None
    assert detect_kind(b"hello") is None


def test_pdf_passthrough_and_rejections() -> None:
    assert to_pdf(b"%PDF-1.4 x", "a.pdf", "soffice", 5) == b"%PDF-1.4 x"
    with pytest.raises(ApiError) as exc:
        to_pdf(b"not a deck", "notes.txt", "soffice", 5)
    assert exc.value.status_code == 400 and exc.value.code == "INVALID_FILE"
    with pytest.raises(ApiError) as exc:
        to_pdf(_zip({"ppt/presentation.xml": b"<p/>"}), "a.pptx", "no-such-soffice-binary", 5)
    assert exc.value.status_code == 415 and exc.value.code == "CONVERTER_MISSING"


@pytest.mark.skipif(not HAS_SOFFICE, reason="LibreOffice not installed")
def test_pptx_converts_to_pdf_with_one_page_per_slide(tmp_path: Path) -> None:
    pdf = to_pdf(_pptx(["Eye Aspect Ratio", "PERCLOS"]), "deck.pptx", "soffice", 120)
    assert pdf.startswith(b"%PDF")
    with fitz.open(stream=pdf, filetype="pdf") as doc:
        assert doc.page_count == 2
        assert "PERCLOS" in doc[1].get_text()


@pytest.mark.skipif(not HAS_SOFFICE, reason="LibreOffice not installed")
def test_upload_pptx_through_api(client: TestClient) -> None:
    itoken, _ = register_login(client, "instructor", "ppt.teacher@uni.edu")
    cls = client.post(
        "/api/v1/classes", json={"name": "Slides", "course_code": "CS-1"}, headers=auth(itoken)
    ).json()
    res = client.post(
        f"/api/v1/classes/{cls['id']}/slides",
        files={"file": ("week1.pptx", _pptx(["Head pose", "Privacy by design"]))},
        headers=auth(itoken),
    )
    assert res.status_code == 200, res.text
    assert res.json()["slide_count"] == 2
    listed = client.get(f"/api/v1/classes/{cls['id']}/slides", headers=auth(itoken)).json()
    assert "Privacy by design" in listed[1]["text_preview"]
    bad = client.post(
        f"/api/v1/classes/{cls['id']}/slides",
        files={"file": ("notes.txt", b"plain text")},
        headers=auth(itoken),
    )
    assert bad.status_code == 400 and bad.json()["error"]["code"] == "INVALID_FILE"
